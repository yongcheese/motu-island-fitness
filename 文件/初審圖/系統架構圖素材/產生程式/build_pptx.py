import re, json
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

PX = 9525  # 1px = 9525 EMU（96 dpi）
W, H = 1900, 1000
svg = open('figma_layers.svg', encoding='utf-8').read()
order = re.findall(r'<(?:g|svg) id="([^"]+)"', svg)
order = [i for i in order if i != 'labels']
labels = json.load(open('elements4/_labels.json', encoding='utf-8'))

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(W * PX), Emu(H * PX)
slide = prs.slides.add_slide(prs.slide_layouts[6])

def bbox(el):
    s = open(f'elements4/{el}.svg', encoding='utf-8').read()
    if el == 'user-coach-aji':
        return (-4, 368, 250, 258)
    x, y, w, h = map(float, re.search(r'viewBox="([^"]+)"', s).group(1).split())
    return x, y, w, h

for el in order:
    x, y, w, h = bbox(el)
    pic = slide.shapes.add_picture(f'png/{el}.png', Emu(int(x * PX)), Emu(int(y * PX)), Emu(int(w * PX)), Emu(int(h * PX)))
    pic.name = el

for t in labels:
    size = t['size']; s = t['s']
    wpx = max(40, int(sum(size if ord(c) > 255 else size * 0.58 for c in s) + 24))
    hpx = int(size * 1.4)
    left = t['x'] if t['anchor'] == 'start' else t['x'] - wpx / 2
    top = t['y'] - size * 1.05
    tb = slide.shapes.add_textbox(Emu(int(left * PX)), Emu(int(top * PX)), Emu(wpx * PX), Emu(hpx * PX))
    tb.name = t['id']
    tf = tb.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT if t['anchor'] == 'start' else PP_ALIGN.CENTER
    r = p.add_run(); r.text = s
    r.font.size = Pt(size * 0.75)
    r.font.name = 'Huninn'
    c = t['color'].lstrip('#'); r.font.color.rgb = RGBColor.from_string(c.upper())

prs.save('系統架構圖_Google簡報.pptx')
print('ok', len(order), 'images', len(labels), 'text boxes')
