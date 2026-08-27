#!/usr/bin/env python3
"""簡易 Markdown → Word 轉檔器（樣式從簡，供內部閱讀用）。

與同目錄的 build_docx.py 不同：這支不做封面、不做目錄、不渲染 mermaid，
只求乾淨好讀。任何 md 都能轉。

用法：
    python3 build_simple.py <來源.md> [輸出.docx]

輸出省略時，會在來源 md 旁邊產生同名的 .docx。
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CN = '蘋方-繁'          # 中文字型
EN = 'Helvetica Neue'   # 西文字型
MONO = 'Menlo'          # 等寬（程式碼）
INK = RGBColor(0x22, 0x22, 0x22)
BLUE = RGBColor(0x1F, 0x4E, 0x79)
GREY = RGBColor(0x66, 0x66, 0x66)


def set_font(run, cn=CN, en=EN):
    """python-docx 不會自動設東亞字型，要直接寫 rPr。"""
    run.font.name = en
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = rpr.makeelement(qn('w:rFonts'), {})
        rpr.append(rfonts)
    rfonts.set(qn('w:ascii'), en)
    rfonts.set(qn('w:hAnsi'), en)
    rfonts.set(qn('w:eastAsia'), cn)


def add_inline(para, text, size=Pt(11), color=INK, bold=False):
    """處理 **粗體** 與 `行內程式碼`，其餘照原樣輸出。"""
    for part in re.split(r'(\*\*.+?\*\*|`[^`]+`)', text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = para.add_run(part[2:-2])
            run.bold = True
            set_font(run)
        elif part.startswith('`') and part.endswith('`'):
            run = para.add_run(part[1:-1])
            set_font(run, cn=MONO, en=MONO)
            run.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
        else:
            run = para.add_run(part)
            run.bold = bold
            set_font(run)
        run.font.size = size
        if not (part.startswith('`') and part.endswith('`')):
            run.font.color.rgb = color


def setup(doc):
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    sec.left_margin = sec.right_margin = Cm(2.2)
    style = doc.styles['Normal']
    style.font.size = Pt(11)
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(6)


def heading(doc, level, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    sizes = {1: Pt(20), 2: Pt(15), 3: Pt(12.5), 4: Pt(11.5)}
    pf.space_before = Pt(4 if level == 1 else 14)
    pf.space_after = Pt(8 if level == 1 else 6)
    pf.keep_with_next = True
    add_inline(p, text, size=sizes.get(level, Pt(11)),
               color=BLUE if level <= 2 else INK, bold=True)
    return p


def code_block(doc, lines):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.5)
    pf.space_before = pf.space_after = Pt(6)
    pf.line_spacing = 1.2
    for i, ln in enumerate(lines):
        run = p.add_run(ln)
        set_font(run, cn=MONO, en=MONO)
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        if i < len(lines) - 1:
            run.add_break()


def quote(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    add_inline(p, text, size=Pt(10.5), color=GREY)


def split_row(row):
    return [c.strip() for c in row.strip().strip('|').split('|')]


def make_table(doc, rows):
    header, body = rows[0], rows[2:]   # rows[1] 是 |---| 分隔列
    t = doc.add_table(rows=1, cols=len(header))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, txt in zip(t.rows[0].cells, header):
        cell.text = ''
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        add_inline(p, txt, size=Pt(10), bold=True, color=BLUE)
    for r in body:
        cells = t.add_row().cells
        for cell, txt in zip(cells, r):
            cell.text = ''
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, txt.replace('<br>', ' '), size=Pt(10))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def convert(src: Path, out: Path):
    doc = Document()
    setup(doc)
    lines = src.read_text(encoding='utf-8').split('\n')
    i = 0
    while i < len(lines):
        ln = lines[i]
        stripped = ln.strip()

        # 程式碼區塊
        if stripped.startswith('```'):
            buf = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                buf.append(lines[i])
                i += 1
            code_block(doc, buf)
            i += 1
            continue

        # 表格：連續的 | 開頭列
        if stripped.startswith('|'):
            buf = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                buf.append(split_row(lines[i]))
                i += 1
            if len(buf) >= 2:
                make_table(doc, buf)
            continue

        # 標題
        m = re.match(r'^(#{1,4})\s+(.*)', stripped)
        if m:
            heading(doc, len(m.group(1)), m.group(2))
            i += 1
            continue

        # 分隔線
        if stripped in ('---', '***', '___'):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            i += 1
            continue

        # 引言
        if stripped.startswith('>'):
            quote(doc, stripped.lstrip('> ').strip())
            i += 1
            continue

        # 清單
        m = re.match(r'^([-*]|\d+\.)\s+(.*)', stripped)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, ('• ' if m.group(1) in '-*' else m.group(1) + ' ') + m.group(2))
            i += 1
            continue

        # 空行
        if not stripped:
            i += 1
            continue

        # 一般段落
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_inline(p, stripped)
        i += 1

    doc.save(out)
    return out


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    src = Path(sys.argv[1]).resolve()
    out = Path(sys.argv[2]).resolve() if len(sys.argv) > 2 else src.with_suffix('.docx')
    convert(src, out)
    print(f'完成：{out}')
