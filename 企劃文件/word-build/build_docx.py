# -*- coding: utf-8 -*-
"""把整合版企劃書 Markdown 轉成美化排版的 Word 檔（含 mermaid 圖轉 PNG）。"""
import re, struct, pathlib
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import subprocess

BASE = pathlib.Path(__file__).parent
SRC_MD = BASE.parent / "AI健身系統_育秀盃企劃書_v2_整合版.md"
OUT = BASE.parent / "AI健身系統_育秀盃企劃書_v2_整合版.docx"
MMD_DIR = BASE / "mmd"
IMG_DIR = BASE / "img"

CN = "PingFang TC"           # 中文字型
EN = "Helvetica Neue"        # 西文字型
NAVY = RGBColor(0x1B, 0x4F, 0x72)
BLUE = RGBColor(0x2C, 0x5F, 0x8D)
SLATE = RGBColor(0x33, 0x3F, 0x4A)
GREY = RGBColor(0x6B, 0x7A, 0x88)
CONTENT_W_CM = 16.0

# ---------- 低階 XML 工具 ----------

# OOXML 規定各 *Pr 的子元素必須依固定順序排列，順序錯誤 Word 會判定檔案毀損。
PPR_ORDER = ['w:pStyle', 'w:keepNext', 'w:keepLines', 'w:pageBreakBefore', 'w:framePr',
             'w:widowControl', 'w:numPr', 'w:suppressLineNumbers', 'w:pBdr', 'w:shd', 'w:tabs',
             'w:suppressAutoHyphens', 'w:kinsoku', 'w:wordWrap', 'w:overflowPunct',
             'w:topLinePunct', 'w:autoSpaceDE', 'w:autoSpaceDN', 'w:bidi', 'w:adjustRightInd',
             'w:snapToGrid', 'w:spacing', 'w:ind', 'w:contextualSpacing', 'w:mirrorIndents',
             'w:suppressOverlap', 'w:jc', 'w:textDirection', 'w:textAlignment',
             'w:textboxTightWrap', 'w:outlineLvl', 'w:divId', 'w:cnfStyle', 'w:rPr',
             'w:sectPr', 'w:pPrChange']
TCPR_ORDER = ['w:cnfStyle', 'w:tcW', 'w:gridSpan', 'w:hMerge', 'w:vMerge', 'w:tcBorders',
              'w:shd', 'w:noWrap', 'w:tcMar', 'w:textDirection', 'w:tcFitText', 'w:vAlign',
              'w:hideMark', 'w:headers']
TBLPR_ORDER = ['w:tblStyle', 'w:tblpPr', 'w:tblOverlap', 'w:bidiVisual', 'w:tblStyleRowBandSize',
               'w:tblStyleColBandSize', 'w:tblW', 'w:jc', 'w:tblCellSpacing', 'w:tblInd',
               'w:tblBorders', 'w:shd', 'w:tblLayout', 'w:tblCellMar', 'w:tblLook',
               'w:tblCaption', 'w:tblDescription']
TRPR_ORDER = ['w:cnfStyle', 'w:divId', 'w:gridBefore', 'w:gridAfter', 'w:wBefore', 'w:wAfter',
              'w:cantSplit', 'w:trHeight', 'w:tblHeader', 'w:tblCellSpacing', 'w:jc', 'w:hidden']


def put(parent, tag, order, replace=True):
    """依 schema 順序插入（或取回既有的）子元素。"""
    existing = parent.find(qn(tag))
    if existing is not None:
        if not replace:
            return existing
        parent.remove(existing)
    el = OxmlElement(tag)
    idx = order.index(tag)
    for child in parent.iterchildren():
        ctag = child.tag.split('}')[-1]
        cq = 'w:' + ctag
        if cq in order and order.index(cq) > idx:
            child.addprevious(el)
            return el
    parent.append(el)
    return el


def set_font(run, cn=CN, en=EN):
    run.font.name = en          # python-docx 會自行建立位置正確的 rFonts
    rf = run._element.get_or_add_rPr().get_or_add_rFonts()
    rf.set(qn('w:ascii'), en); rf.set(qn('w:hAnsi'), en)
    rf.set(qn('w:eastAsia'), cn); rf.set(qn('w:cs'), en)


def para_border(p, edges):
    """edges: {'bottom': (size8th, color)} 之類。"""
    pPr = p._p.get_or_add_pPr()
    bdr = put(pPr, 'w:pBdr', PPR_ORDER)
    for edge in ('top', 'left', 'bottom', 'right'):
        if edge not in edges:
            continue
        sz, color = edges[edge]
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz))
        e.set(qn('w:space'), '4'); e.set(qn('w:color'), color)
        bdr.append(e)


def para_shade(p, color):
    sh = put(p._p.get_or_add_pPr(), 'w:shd', PPR_ORDER)
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), color)


def cell_shade(cell, color):
    sh = put(cell._tc.get_or_add_tcPr(), 'w:shd', TCPR_ORDER)
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), color)


def cell_margins(cell, top=60, bottom=60, left=110, right=110):
    mar = put(cell._tc.get_or_add_tcPr(), 'w:tcMar', TCPR_ORDER)
    for tag, val in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        e = OxmlElement('w:' + tag)
        e.set(qn('w:w'), str(val)); e.set(qn('w:type'), 'dxa')
        mar.append(e)


def cell_borders(cell, spec):
    """spec: {'left': (sz, color), ...}"""
    borders = put(cell._tc.get_or_add_tcPr(), 'w:tcBorders', TCPR_ORDER)
    for edge in ('top', 'left', 'bottom', 'right'):
        if edge not in spec:
            continue
        sz, color = spec[edge]
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz))
        e.set(qn('w:space'), '0'); e.set(qn('w:color'), color)
        borders.append(e)


def table_borders(table, color='C9D6E2', sz=4, inside_h='DCE5EE'):
    borders = put(table._tbl.tblPr, 'w:tblBorders', TBLPR_ORDER)
    for tag, c in (('top', color), ('left', color), ('bottom', color), ('right', color),
                   ('insideH', inside_h), ('insideV', inside_h)):
        e = OxmlElement('w:' + tag)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz))
        e.set(qn('w:space'), '0'); e.set(qn('w:color'), c)
        borders.append(e)


def table_full_width(table):
    """表格寬度設為版心 100%，並允許 Word 自動配欄寬。"""
    tblPr = table._tbl.tblPr
    w = put(tblPr, 'w:tblW', TBLPR_ORDER)
    w.set(qn('w:w'), '5000'); w.set(qn('w:type'), 'pct')
    layout = put(tblPr, 'w:tblLayout', TBLPR_ORDER)
    layout.set(qn('w:type'), 'autofit')


def repeat_header(row):
    e = put(row._tr.get_or_add_trPr(), 'w:tblHeader', TRPR_ORDER)
    e.set(qn('w:val'), 'true')


def keep_with_next(p):
    e = put(p._p.get_or_add_pPr(), 'w:keepNext', PPR_ORDER)
    e.set(qn('w:val'), 'true')


def add_field(paragraph, instr):
    r = paragraph.add_run()
    fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr
    fc2 = OxmlElement('w:fldChar'); fc2.set(qn('w:fldCharType'), 'separate')
    fc3 = OxmlElement('w:fldChar'); fc3.set(qn('w:fldCharType'), 'end')
    r._r.append(fc); r._r.append(it); r._r.append(fc2); r._r.append(fc3)
    return r


def png_size(path):
    with open(path, 'rb') as f:
        head = f.read(26)
    w, h = struct.unpack('>II', head[16:24])
    return w, h



# ---------- 前處理：抽出 mermaid 並轉成 PNG ----------

def prepare():
    """把 mermaid 區塊抽成 .mmd 檔並渲染成 PNG，回傳去掉 mermaid 的正文行。"""
    MMD_DIR.mkdir(exist_ok=True); IMG_DIR.mkdir(exist_ok=True)
    out, buf, inblk, idx = [], [], False, 0
    for ln in SRC_MD.read_text(encoding='utf-8').split('\n'):
        if not inblk and ln.strip().startswith('```mermaid'):
            inblk, buf = True, []; continue
        if inblk and ln.strip() == '```':
            inblk = False; idx += 1
            mmd = MMD_DIR / f'diagram{idx:02d}.mmd'
            png = IMG_DIR / f'diagram{idx:02d}.png'
            src = '\n'.join(buf) + '\n'
            if not mmd.exists() or mmd.read_text(encoding='utf-8') != src:
                mmd.write_text(src, encoding='utf-8')
            if not png.exists() or png.stat().st_mtime < mmd.stat().st_mtime:
                render(mmd, png)
            out.append(f'![](img/diagram{idx:02d}.png)')
            continue
        if inblk:
            buf.append(ln); continue
        out.append(ln)
    print(f'mermaid 圖 {idx} 張')
    return out


def render(mmd, png):
    """用 mermaid-cli 產圖（需要 Node.js，第一次會自動下載）。"""
    print('  渲染', png.name)
    subprocess.run(['npx', '--yes', '@mermaid-js/mermaid-cli@11',
                    '-i', str(mmd), '-o', str(png),
                    '-c', str(BASE / 'mconf.json'), '-p', str(BASE / 'pconf.json'),
                    '-b', 'white', '-s', '3'], check=True)


# ---------- 行內語法 ----------
INLINE = re.compile(r'(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|https?://\S+)')


def add_inline(paragraph, text, size=Pt(11), color=SLATE, bold=False, italic=False, cn=CN, en=EN):
    text = text.replace('\\', '')
    for part in INLINE.split(text):
        if not part:
            continue
        b, i, code, url = bold, italic, False, None
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            part, b = part[2:-2], True
        elif part.startswith('`') and part.endswith('`'):
            part, code = part[1:-1], True
        elif part.startswith('[') and '](' in part:
            m = re.match(r'\[([^\]]+)\]\(([^)]+)\)', part)
            part, url = m.group(1), m.group(2)
        elif part.startswith('http'):
            url = part
        run = paragraph.add_run(part)
        run.font.size = Pt(size.pt - 0.5) if code else size
        run.bold = b
        run.italic = i
        if code:
            set_font(run, cn='Menlo', en='Menlo')
            run.font.color.rgb = RGBColor(0xA6, 0x2B, 0x3C)
        else:
            set_font(run, cn=cn, en=en)
            run.font.color.rgb = color
        if url:
            run.font.color.rgb = RGBColor(0x1F, 0x5C, 0x99)
            run.font.underline = True
    return paragraph


# ---------- 解析 ----------

def parse_blocks(lines):
    blocks, i, n = [], 0, len(lines)
    while i < n:
        ln = lines[i]
        s = ln.strip()
        if not s:
            i += 1; continue
        if s in ('---', '***', '___'):
            blocks.append(('hr', None)); i += 1; continue
        m = re.match(r'^(#{1,6})\s+(.*)$', s)
        if m:
            blocks.append(('h', (len(m.group(1)), m.group(2).strip()))); i += 1; continue
        m = re.match(r'^!\[[^\]]*\]\(([^)]+)\)\s*$', s)
        if m:
            blocks.append(('img', m.group(1))); i += 1; continue
        if s.startswith('>'):
            buf = []
            while i < n and lines[i].strip().startswith('>'):
                buf.append(re.sub(r'^\s*>\s?', '', lines[i]).rstrip()); i += 1
            blocks.append(('quote', buf)); continue
        if s.startswith('|') and i + 1 < n and re.match(r'^\|[\s:*\-|]+\|$', lines[i + 1].strip()):
            rows = []
            while i < n and lines[i].strip().startswith('|'):
                rows.append(lines[i].strip()); i += 1
            blocks.append(('table', rows)); continue
        if re.match(r'^(\-|\*)\s+', s) or re.match(r'^\d+\.\s+', s):
            items = []
            while i < n and lines[i].strip() and (re.match(r'^(\-|\*)\s+', lines[i].strip())
                                                 or re.match(r'^\d+\.\s+', lines[i].strip())):
                t = lines[i].strip()
                m2 = re.match(r'^(\d+\.|\-|\*)\s+(.*)$', t)
                items.append((m2.group(1), m2.group(2)))
                i += 1
            blocks.append(('list', items)); continue
        buf = []
        while i < n and lines[i].strip() and not re.match(r'^(#{1,6}\s|\||>|!\[|\-\s|\*\s|\d+\.\s|---$)', lines[i].strip()):
            buf.append(lines[i].strip()); i += 1
        if buf:
            blocks.append(('p', ' '.join(buf)))
        else:
            i += 1
    return blocks


def split_row(row):
    row = row.strip()
    if row.startswith('|'):
        row = row[1:]
    if row.endswith('|'):
        row = row[:-1]
    cells, cur, esc = [], '', False
    for ch in row:
        if esc:
            cur += ch; esc = False; continue
        if ch == '\\':
            esc = True; continue
        if ch == '|':
            cells.append(cur.strip()); cur = ''
        else:
            cur += ch
    cells.append(cur.strip())
    return cells


# ---------- 樣式 ----------

def setup_styles(doc):
    st = doc.styles
    normal = st['Normal']
    normal.font.size = Pt(11)
    normal.font.color.rgb = SLATE
    normal.font.name = EN
    rf = normal.element.rPr.rFonts
    rf.set(qn('w:eastAsia'), CN); rf.set(qn('w:ascii'), EN); rf.set(qn('w:hAnsi'), EN)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.4
    pf.space_after = Pt(6)
    pf.space_before = Pt(0)

    specs = {
        'Heading 1': (17, True, NAVY, 20, 8),
        'Heading 2': (14, True, BLUE, 14, 6),
        'Heading 3': (12.5, True, RGBColor(0x2F, 0x40, 0x52), 10, 4),
        'Heading 4': (11.5, True, RGBColor(0x44, 0x55, 0x66), 8, 3),
    }
    for name, (size, bold, color, before, after) in specs.items():
        s = st[name]
        s.font.size = Pt(size); s.font.bold = bold; s.font.color.rgb = color
        s.font.name = EN
        rfn = s.element.get_or_add_rPr().get_or_add_rFonts()
        rfn.set(qn('w:eastAsia'), CN); rfn.set(qn('w:ascii'), EN); rfn.set(qn('w:hAnsi'), EN)
        p = s.paragraph_format
        p.space_before = Pt(before); p.space_after = Pt(after)
        p.line_spacing_rule = WD_LINE_SPACING.MULTIPLE; p.line_spacing = 1.2
        p.keep_with_next = True


def setup_page(section):
    section.page_width = Cm(21); section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5); section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.2); section.bottom_margin = Cm(2.2)
    section.header_distance = Cm(1.2); section.footer_distance = Cm(1.2)


def build_cover(doc, title, subtitle, meta_lines):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('育秀盃創意獎　競賽企劃書'); r.font.size = Pt(13); r.bold = True
    r.font.color.rgb = GREY; set_font(r)
    p.paragraph_format.space_after = Pt(4)

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('AI 應用類'); r.font.size = Pt(11); r.font.color.rgb = GREY; set_font(r)
    p.paragraph_format.space_after = Pt(26)

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Island Fitness AI'); r.font.size = Pt(30); r.bold = True
    r.font.color.rgb = NAVY; set_font(r)
    p.paragraph_format.space_after = Pt(6)

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('AI 健身動作品質判定與島嶼建設連動系統')
    r.font.size = Pt(16); r.bold = True; r.font.color.rgb = BLUE; set_font(r)
    p.paragraph_format.space_after = Pt(14)

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para_border(p, {'bottom': (12, '2C5F8D')})
    p.paragraph_format.space_after = Pt(20)

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('不是計次，而是品質確認')
    r.font.size = Pt(12.5); r.font.color.rgb = SLATE; set_font(r)
    p.paragraph_format.space_after = Pt(34)

    # 資訊表
    rows = []
    for ln in meta_lines:
        m = re.match(r'^\*\*(.+?)\*\*[:：]\s*(.*)$', ln.strip())
        if m:
            rows.append((m.group(1), m.group(2)))
        elif ln.strip():
            rows.append(('', ln.strip()))
    t = doc.add_table(rows=len(rows), cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, (k, v) in enumerate(rows):
        c0, c1 = t.rows[i].cells
        c0.width = Cm(4.2); c1.width = Cm(10.8)
        p0 = c0.paragraphs[0]; p0.paragraph_format.space_after = Pt(3)
        r0 = p0.add_run(k); r0.bold = True; r0.font.size = Pt(10.5)
        r0.font.color.rgb = BLUE; set_font(r0)
        p1 = c1.paragraphs[0]; p1.paragraph_format.space_after = Pt(3)
        add_inline(p1, v, size=Pt(10.5))
        cell_margins(c0, 70, 70, 60, 60); cell_margins(c1, 70, 70, 60, 60)
    table_borders(t, color='FFFFFF', sz=2, inside_h='E6EDF4')

    doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('國立臺中科技大學')
    r.font.size = Pt(11); r.bold = True; r.font.color.rgb = GREY; set_font(r)

    doc.add_page_break()


def build_toc(doc):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('目　　錄'); r.font.size = Pt(18); r.bold = True
    r.font.color.rgb = NAVY; set_font(r)
    p.paragraph_format.space_after = Pt(6)
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para_border(p2, {'bottom': (8, '2C5F8D')})
    p2.paragraph_format.space_after = Pt(14)

    p3 = doc.add_paragraph()
    add_field(p3, ' TOC \\o "1-2" \\h \\z \\u ')
    hint = doc.add_paragraph()
    r = hint.add_run('（若目錄未顯示頁碼，請在此處按右鍵選「更新功能變數」）')
    r.font.size = Pt(9); r.font.color.rgb = GREY; r.italic = True; set_font(r)
    hint.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()


def add_footer(doc):
    sec = doc.sections[0]
    f = sec.footer.paragraphs[0]
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = f.add_run('第 '); r.font.size = Pt(9); r.font.color.rgb = GREY; set_font(r)
    add_field(f, ' PAGE ')
    for run in f.runs:
        run.font.size = Pt(9); run.font.color.rgb = GREY; set_font(run)
    r = f.add_run(' 頁'); r.font.size = Pt(9); r.font.color.rgb = GREY; set_font(r)

    h = sec.header.paragraphs[0]
    h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = h.add_run('Island Fitness AI ── AI 健身動作品質判定與島嶼建設連動系統')
    r.font.size = Pt(8.5); r.font.color.rgb = GREY; set_font(r)
    para_border(h, {'bottom': (4, 'C9D6E2')})


SETTINGS_ORDER = ['w:writeProtection', 'w:view', 'w:zoom', 'w:removePersonalInformation',
                  'w:removeDateAndTime', 'w:doNotDisplayPageBoundaries', 'w:displayBackgroundShape',
                  'w:printPostScriptOverText', 'w:printFractionalCharacterWidth',
                  'w:printFormsData', 'w:embedTrueTypeFonts', 'w:embedSystemFonts',
                  'w:saveSubsetFonts', 'w:saveFormsData', 'w:mirrorMargins', 'w:alignBordersAndEdges',
                  'w:bordersDoNotSurroundHeader', 'w:bordersDoNotSurroundFooter', 'w:gutterAtTop',
                  'w:hideSpellingErrors', 'w:hideGrammaticalErrors', 'w:activeWritingStyle',
                  'w:proofState', 'w:formsDesign', 'w:attachedTemplate', 'w:linkStyles',
                  'w:stylePaneFormatFilter', 'w:stylePaneSortMethod', 'w:documentType',
                  'w:mailMerge', 'w:revisionView', 'w:trackChanges', 'w:documentProtection',
                  'w:autoFormatOverride', 'w:styleLockTheme', 'w:styleLockQFSet',
                  'w:defaultTabStop', 'w:autoHyphenation', 'w:consecutiveHyphenLimit',
                  'w:hyphenationZone', 'w:doNotHyphenateCaps', 'w:showEnvelope',
                  'w:summaryLength', 'w:clickAndTypeStyle', 'w:defaultTableStyle',
                  'w:evenAndOddHeaders', 'w:bookFoldRevPrinting', 'w:bookFoldPrinting',
                  'w:bookFoldPrintingSheets', 'w:drawingGridHorizontalSpacing',
                  'w:drawingGridVerticalSpacing', 'w:displayHorizontalDrawingGridEvery',
                  'w:displayVerticalDrawingGridEvery', 'w:doNotUseMarginsForDrawingGridOrigin',
                  'w:drawingGridHorizontalOrigin', 'w:drawingGridVerticalOrigin',
                  'w:doNotShadeFormData', 'w:noPunctuationKerning', 'w:characterSpacingControl',
                  'w:printTwoOnOne', 'w:strictFirstAndLastChars', 'w:noLineBreaksAfter',
                  'w:noLineBreaksBefore', 'w:savePreviewPicture', 'w:doNotValidateAgainstSchema',
                  'w:saveInvalidXml', 'w:ignoreMixedContent', 'w:alwaysShowPlaceholderText',
                  'w:doNotDemarcateInvalidXml', 'w:saveXmlDataOnly', 'w:useXSLTWhenSaving',
                  'w:saveThroughXslt', 'w:showXMLTags', 'w:alwaysMergeEmptyNamespace',
                  'w:updateFields', 'w:hdrShapeDefaults', 'w:footnotePr', 'w:endnotePr',
                  'w:compat', 'w:docVars', 'w:rsids', 'm:mathPr', 'w:attachedSchema',
                  'w:themeFontLang', 'w:clrSchemeMapping', 'w:doNotIncludeSubdocsInStats',
                  'w:doNotAutoCompressPictures', 'w:forceUpgrade', 'w:captions',
                  'w:readModeInkLockDown', 'w:smartTagType', 'w:shapeDefaults',
                  'w:doNotEmbedSmartTags', 'w:decimalSymbol', 'w:listSeparator']


def update_fields_on_open(doc):
    """開檔時提示更新功能變數，目錄頁碼才會自動填上。"""
    e = put(doc.settings.element, 'w:updateFields', SETTINGS_ORDER)
    e.set(qn('w:val'), 'true')


# ---------- 區塊寫入 ----------

def write_heading(doc, level, text):
    lvl = min(max(level - 1, 1), 4)   # ## -> H1
    p = doc.add_heading(level=lvl)
    color = {1: NAVY, 2: BLUE, 3: RGBColor(0x2F, 0x40, 0x52), 4: RGBColor(0x44, 0x55, 0x66)}[lvl]
    size = {1: Pt(17), 2: Pt(14), 3: Pt(12.5), 4: Pt(11.5)}[lvl]
    add_inline(p, text, size=size, color=color, bold=True)
    if lvl == 1:
        para_border(p, {'bottom': (10, '2C5F8D')})
    elif lvl == 2:
        para_border(p, {'left': (18, '9FC0DC')})
        p.paragraph_format.left_indent = Cm(0.25)
    keep_with_next(p)
    return p


def write_paragraph(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    add_inline(p, text)
    return p


def write_quote(doc, lines):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.rows[0].cells[0]
    cell.width = Cm(CONTENT_W_CM)
    cell_shade(cell, 'F2F7FB')
    cell_margins(cell, 120, 120, 200, 160)
    first = True
    for ln in lines:
        if not ln.strip():
            continue
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.35
        add_inline(p, ln.strip(), size=Pt(10.5), color=RGBColor(0x24, 0x3A, 0x4C))
    # 左側粗藍條
    cell_borders(cell, {'top': (2, 'DCE7F1'), 'left': (24, '2C5F8D'),
                        'bottom': (2, 'DCE7F1'), 'right': (2, 'DCE7F1')})
    table_full_width(t)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def write_list(doc, items):
    """自行輸出項目符號／編號，確保與原文逐字相同（不用 Word 自動編號）。"""
    for marker_src, text in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.35
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        marker = '• ' if marker_src in ('-', '*') else marker_src + ' '
        r = p.add_run(marker)
        r.font.size = Pt(11); r.font.color.rgb = BLUE; r.bold = True
        set_font(r)
        add_inline(p, text)


def write_image(doc, path):
    full = BASE / path
    w, h = png_size(full)
    width_cm = min(CONTENT_W_CM, 15.5)
    height_cm = width_cm * h / w
    max_h = 19.5
    if height_cm > max_h:
        height_cm = max_h
        width_cm = height_cm * w / h
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(10)
    p.add_run().add_picture(str(full), width=Cm(width_cm))


def write_table(doc, rows):
    header = split_row(rows[0])
    aligns = split_row(rows[1])
    data = [split_row(r) for r in rows[2:]]
    ncol = len(header)
    data = [d + [''] * (ncol - len(d)) if len(d) < ncol else d[:ncol] for d in data]

    def al(spec):
        spec = spec.strip()
        if spec.startswith(':') and spec.endswith(':'):
            return WD_ALIGN_PARAGRAPH.CENTER
        if spec.endswith(':'):
            return WD_ALIGN_PARAGRAPH.RIGHT
        return WD_ALIGN_PARAGRAPH.LEFT

    alignments = [al(a) for a in aligns] + [WD_ALIGN_PARAGRAPH.LEFT] * ncol

    t = doc.add_table(rows=len(data) + 1, cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    # 表頭
    for j, htxt in enumerate(header):
        c = t.rows[0].cells[j]
        cell_shade(c, '2C5F8D')
        cell_margins(c)
        p = c.paragraphs[0]
        p.alignment = alignments[j]
        p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.2
        add_inline(p, htxt, size=Pt(10), color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)
    repeat_header(t.rows[0])
    for i, row in enumerate(data):
        for j, txt in enumerate(row):
            c = t.rows[i + 1].cells[j]
            if i % 2 == 1:
                cell_shade(c, 'F5F8FB')
            cell_margins(c)
            p = c.paragraphs[0]
            p.alignment = alignments[j]
            p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.25
            add_inline(p, txt, size=Pt(10))
    table_borders(t)
    table_full_width(t)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def write_hr(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(10)
    para_border(p, {'bottom': (6, 'D8E1EA')})


# ---------- 主流程 ----------

def main():
    lines = prepare()
    hr_idx = next(i for i, l in enumerate(lines) if l.strip() == '---')
    head, rest = lines[:hr_idx], lines[hr_idx + 1:]

    title = next(l[2:].strip() for l in head if l.startswith('# '))
    subtitle = next(l[3:].strip() for l in head if l.startswith('## '))
    meta = [re.sub(r'^\s*>\s?', '', l).strip() for l in head if l.strip().startswith('>')]

    doc = Document()
    setup_styles(doc)
    setup_page(doc.sections[0])
    add_footer(doc)
    update_fields_on_open(doc)

    build_cover(doc, title, subtitle, meta)
    build_toc(doc)

    for kind, payload in parse_blocks(rest):
        if kind == 'h':
            write_heading(doc, payload[0], payload[1])
        elif kind == 'p':
            write_paragraph(doc, payload)
        elif kind == 'quote':
            write_quote(doc, payload)
        elif kind == 'list':
            write_list(doc, payload)
        elif kind == 'table':
            write_table(doc, payload)
        elif kind == 'img':
            write_image(doc, payload)
        elif kind == 'hr':
            write_hr(doc)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print('saved:', OUT)


if __name__ == '__main__':
    main()
